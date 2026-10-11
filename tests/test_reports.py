import io
import openpyxl
from app.services.reports import compute_kpis, generate_placement_excel

def test_compute_kpis_empty():
    kpis = compute_kpis([])
    assert kpis['total_registered'] == 0
    assert kpis['total_placed'] == 0
    assert kpis['placement_percentage'] == 0.0
    assert kpis['highest_package'] == 0.0
    assert kpis['average_package'] == 0.0
    assert kpis['median_package'] == 0.0


def test_compute_kpis_standard():
    student_records = [
        {'name': 'A', 'is_placed': True, 'package': 10.0},
        {'name': 'B', 'is_placed': True, 'package': 20.0},
        {'name': 'C', 'is_placed': True, 'package': 30.0},
        {'name': 'D', 'is_placed': False, 'package': 0.0},
    ]
    kpis = compute_kpis(student_records)
    assert kpis['total_registered'] == 4
    assert kpis['total_placed'] == 3
    assert kpis['placement_percentage'] == 75.0
    assert kpis['highest_package'] == 30.0
    assert kpis['average_package'] == 20.0
    assert kpis['median_package'] == 20.0


def test_compute_kpis_zero_placed():
    student_records = [
        {'name': 'A', 'is_placed': False},
        {'name': 'B', 'is_placed': False},
    ]
    kpis = compute_kpis(student_records)
    assert kpis['total_registered'] == 2
    assert kpis['total_placed'] == 0
    assert kpis['placement_percentage'] == 0.0
    assert kpis['highest_package'] == 0.0


def test_generate_student_wise_excel():
    raw_data = {
        'students': [
            {
                'roll_no': '22CS0101',
                'name': 'Aarav Sharma',
                'branch': 'CSE',
                'passing_year': 2026,
                'cgpa': 9.2,
                'active_backlogs': 0,
                'is_placed': True,
                'company_name': 'Google',
                'drive_title': 'SDE 1',
                'package': 32.0,
                'offer_date': '2026-10-01'
            },
            {
                'roll_no': '22IT0201',
                'name': 'Kavya Pillai',
                'branch': 'IT',
                'passing_year': 2026,
                'cgpa': 8.5,
                'active_backlogs': 0,
                'is_placed': False,
                'company_name': '-',
                'drive_title': '-',
                'package': 0.0,
                'offer_date': '-'
            }
        ]
    }
    filters = {'Batch_Year': '2026', 'Branch': 'All'}
    excel_stream = generate_placement_excel('student_wise', raw_data, filters)
    assert isinstance(excel_stream, io.BytesIO)

    # Validate workbook structure
    wb = openpyxl.load_workbook(excel_stream)
    assert "Executive Summary" in wb.sheetnames
    assert "Student Details" in wb.sheetnames

    # Check Executive summary title
    sum_sheet = wb["Executive Summary"]
    assert "Training & Placement Office" in str(sum_sheet["B2"].value)

    # Check Student Details sheet header and data
    data_sheet = wb["Student Details"]
    assert data_sheet["A1"].value == "Roll No"
    assert data_sheet["B1"].value == "Student Name"
    assert data_sheet["J1"].value == "Package (LPA)"
    assert data_sheet["A2"].value == "22CS0101"
    assert data_sheet["J2"].value == 32.0


def test_generate_company_and_branch_reports():
    raw_data = {
        'students': [{'name': 'A', 'is_placed': True, 'package': 12.0}],
        'company_summary': [
            {'name': 'Google', 'industry': 'Tech', 'total_drives': 1, 'total_applicants': 10, 'placed_count': 2, 'highest_pkg': 32.0, 'avg_pkg': 28.0}
        ],
        'branch_summary': [
            {'branch': 'CSE', 'registered': 10, 'placed': 8, 'placement_pct': 80.0, 'highest_pkg': 32.0, 'avg_pkg': 18.5, 'median_pkg': 16.0}
        ],
        'year_summary': [
            {'batch_year': 2026, 'registered': 30, 'placed': 20, 'placement_pct': 66.7, 'highest_pkg': 32.0, 'avg_pkg': 15.0, 'median_pkg': 12.0}
        ]
    }

    # Company report
    c_stream = generate_placement_excel('company_wise', raw_data)
    wb_c = openpyxl.load_workbook(c_stream)
    assert "Company Summary" in wb_c.sheetnames

    # Branch report
    b_stream = generate_placement_excel('branch_wise', raw_data)
    wb_b = openpyxl.load_workbook(b_stream)
    assert "Branch Summary" in wb_b.sheetnames

    # Year report
    y_stream = generate_placement_excel('year_wise', raw_data)
    wb_y = openpyxl.load_workbook(y_stream)
    assert "Batch Year Summary" in wb_y.sheetnames
